"""Exportação do Relatório de Fechamento em planilha Excel (.xlsx).

Consome a estrutura calculada de ``Relatorio`` (objeto do domínio) sem recalcular
nenhum valor ou total. Células de quantidade e valor são numéricas reais (int e
float formatado como moeda), permitindo conferência e soma no próprio Excel (Req 8.6).
"""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.dominio import Relatorio, RelatorioGeral

# Cores da identidade visual (mesmos valores de frontend/src/styles/tokens.css)
COR_AZUL_900 = "0B2F58"
COR_GELO_50 = "F8FAFD"
COR_GELO_100 = "F0F4F9"
COR_GELO_200 = "E0E7EE"
COR_GELO_300 = "CDD5DF"
COR_GELO_600 = "56626F"
COR_GELO_900 = "141B24"

# Formato monetário padrão brasileiro para células numéricas do Excel
FORMATO_MOEDA = "R$ #,##0.00"


def gerar_excel(relatorio: Relatorio) -> bytes:
    """Gera o arquivo binário Excel (.xlsx) a partir do relatório já calculado.

    Args:
        relatorio: objeto de domínio contendo linhas, colunas e totais agregados.

    Returns:
        Bytes correspondentes ao arquivo .xlsx gerado.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Fechamento"

    # Estilos reutilizáveis
    fonte_titulo_cliente = Font(name="Inter", size=12, bold=True, color=COR_AZUL_900)
    fonte_periodo = Font(name="Inter", size=10, color=COR_GELO_600)
    fonte_cabecalho = Font(name="Inter", size=10, bold=True, color=COR_GELO_900)
    fonte_dados = Font(name="Inter", size=10, color=COR_GELO_900)
    fonte_totais = Font(name="Inter", size=10, bold=True, color=COR_AZUL_900)

    preenchimento_cabecalho = PatternFill(
        start_color=COR_GELO_100, end_color=COR_GELO_100, fill_type="solid"
    )
    preenchimento_totais = PatternFill(
        start_color=COR_GELO_50, end_color=COR_GELO_50, fill_type="solid"
    )

    borda_fina = Side(style="thin", color=COR_GELO_300)
    borda_cabecalho = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)
    borda_dados = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)
    borda_totais = Border(
        left=borda_fina,
        right=borda_fina,
        top=Side(style="medium", color=COR_AZUL_900),
        bottom=Side(style="double", color=COR_AZUL_900),
    )

    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    alinhamento_direita = Alignment(horizontal="right", vertical="center")

    # 1. Metadados do cabeçalho (Linhas 1 a 3)
    ws.cell(row=1, column=1, value=f"Cliente: {relatorio.cliente_nome}").font = fonte_titulo_cliente
    periodo_texto = (
        f"Período: {relatorio.inicio.strftime('%d/%m/%Y')} a {relatorio.fim.strftime('%d/%m/%Y')}"
    )
    ws.cell(row=2, column=1, value=periodo_texto).font = fonte_periodo

    # 2. Cabeçalho da tabela (Linha 4)
    titulos = ["Data", "Comanda"]
    for col in relatorio.colunas_itens:
        titulos.append(col.nome)
    titulos.extend(["Total de peças", "Total R$"])

    linha_cabecalho = 4
    for idx_col, titulo in enumerate(titulos, start=1):
        celula = ws.cell(row=linha_cabecalho, column=idx_col, value=titulo)
        celula.font = fonte_cabecalho
        celula.fill = preenchimento_cabecalho
        celula.border = borda_cabecalho
        if titulo in ("Data", "Comanda"):
            celula.alignment = alinhamento_centro
        else:
            celula.alignment = alinhamento_direita

    # 3. Linhas de dados (a partir da linha 5)
    linha_atual = linha_cabecalho + 1
    for linha in relatorio.linhas:
        # Data
        c_data = ws.cell(row=linha_atual, column=1, value=linha.data.strftime("%d/%m/%Y"))
        c_data.font = fonte_dados
        c_data.alignment = alinhamento_centro
        c_data.border = borda_dados

        # Comanda
        c_comanda = ws.cell(row=linha_atual, column=2, value=linha.comanda or "")
        c_comanda.font = fonte_dados
        c_comanda.alignment = alinhamento_centro
        c_comanda.border = borda_dados

        # Itens
        col_offset = 3
        for col_item in relatorio.colunas_itens:
            qtd = linha.quantidades.get(col_item.item_id)
            c_item = ws.cell(
                row=linha_atual,
                column=col_offset,
                value=int(qtd) if qtd is not None else None,
            )
            c_item.font = fonte_dados
            c_item.alignment = alinhamento_direita
            c_item.border = borda_dados
            col_offset += 1

        # Total de peças
        c_pecas = ws.cell(row=linha_atual, column=col_offset, value=int(linha.total_pecas))
        c_pecas.font = fonte_dados
        c_pecas.alignment = alinhamento_direita
        c_pecas.border = borda_dados
        col_offset += 1

        # Total R$
        c_valor = ws.cell(row=linha_atual, column=col_offset, value=float(linha.total_valor))
        c_valor.font = fonte_dados
        c_valor.alignment = alinhamento_direita
        c_valor.number_format = FORMATO_MOEDA
        c_valor.border = borda_dados

        linha_atual += 1

    # 4. Linha de Totais do Período
    c_tot_rotulo = ws.cell(row=linha_atual, column=1, value="Totais")
    c_tot_rotulo.font = fonte_totais
    c_tot_rotulo.fill = preenchimento_totais
    c_tot_rotulo.alignment = alinhamento_centro
    c_tot_rotulo.border = borda_totais

    c_tot_comanda = ws.cell(row=linha_atual, column=2, value="")
    c_tot_comanda.font = fonte_totais
    c_tot_comanda.fill = preenchimento_totais
    c_tot_comanda.border = borda_totais

    col_offset = 3
    for col_item in relatorio.colunas_itens:
        tot_qtd = relatorio.totais.por_item.get(col_item.item_id, 0)
        c_tot_item = ws.cell(row=linha_atual, column=col_offset, value=int(tot_qtd))
        c_tot_item.font = fonte_totais
        c_tot_item.fill = preenchimento_totais
        c_tot_item.alignment = alinhamento_direita
        c_tot_item.border = borda_totais
        col_offset += 1

    c_tot_pecas = ws.cell(
        row=linha_atual, column=col_offset, value=int(relatorio.totais.total_pecas)
    )
    c_tot_pecas.font = fonte_totais
    c_tot_pecas.fill = preenchimento_totais
    c_tot_pecas.alignment = alinhamento_direita
    c_tot_pecas.border = borda_totais
    col_offset += 1

    c_tot_valor = ws.cell(
        row=linha_atual, column=col_offset, value=float(relatorio.totais.total_valor)
    )
    c_tot_valor.font = fonte_totais
    c_tot_valor.fill = preenchimento_totais
    c_tot_valor.alignment = alinhamento_direita
    c_tot_valor.number_format = FORMATO_MOEDA
    c_tot_valor.border = borda_totais

    # 5. Ajustar largura das colunas ao conteúdo
    for col in ws.iter_cols(min_row=linha_cabecalho, max_row=linha_atual, max_col=len(titulos)):
        max_tam = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                # Se for valor monetário, estimar tamanho com a formatação
                if cell.number_format == FORMATO_MOEDA:
                    val_str = (
                        f"R$ {cell.value:,.2f}".replace(",", "X")
                        .replace(".", ",")
                        .replace("X", ".")
                    )
                else:
                    val_str = str(cell.value)
                max_tam = max(max_tam, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_tam + 4, 12)

    # 6. Congelar a primeira linha da tabela (linha 5 começa os dados roláveis)
    ws.freeze_panes = "A5"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Relatório geral — todos os clientes (v1.1, Req 5.6)
# ---------------------------------------------------------------------------

# Excel recusa estes caracteres em nome de aba e limita a 31 caracteres
_PROIBIDOS_NA_ABA = re.compile(r"[\[\]:*?/\\]")
LIMITE_NOME_ABA = 31
NOME_ABA_RESUMO = "Resumo"


def nome_de_aba(nome: str, usados: set[str]) -> str:
    """Nome de aba válido e único (comparação sem caixa, como o Excel faz).

    Remove os caracteres proibidos, corta em 31 e, se repetir, acrescenta
    " (2)", " (3)"… sem ultrapassar o limite.
    """
    base = _PROIBIDOS_NA_ABA.sub(" ", nome).strip().strip("'") or "Cliente"
    base = base[:LIMITE_NOME_ABA].rstrip()
    candidato = base
    contador = 2
    while candidato.casefold() in usados:
        sufixo = f" ({contador})"
        candidato = base[: LIMITE_NOME_ABA - len(sufixo)].rstrip() + sufixo
        contador += 1
    usados.add(candidato.casefold())
    return candidato


class _Estilos:
    """Estilos do relatório geral, os mesmos do relatório por cliente."""

    titulo = Font(name="Inter", size=12, bold=True, color=COR_AZUL_900)
    periodo = Font(name="Inter", size=10, color=COR_GELO_600)
    cabecalho = Font(name="Inter", size=10, bold=True, color=COR_GELO_900)
    dados = Font(name="Inter", size=10, color=COR_GELO_900)
    totais = Font(name="Inter", size=10, bold=True, color=COR_AZUL_900)
    fundo_cabecalho = PatternFill(
        start_color=COR_GELO_100, end_color=COR_GELO_100, fill_type="solid"
    )
    fundo_totais = PatternFill(start_color=COR_GELO_50, end_color=COR_GELO_50, fill_type="solid")
    _fina = Side(style="thin", color=COR_GELO_300)
    borda = Border(left=_fina, right=_fina, top=_fina, bottom=_fina)
    borda_totais = Border(
        left=_fina,
        right=_fina,
        top=Side(style="medium", color=COR_AZUL_900),
        bottom=Side(style="double", color=COR_AZUL_900),
    )
    esquerda = Alignment(horizontal="left", vertical="center")
    direita = Alignment(horizontal="right", vertical="center")


def _tabela(ws, linha_inicial: int, titulos: list[str], linhas: list[list], total: list) -> None:  # noqa: ANN001
    """Escreve cabeçalho, linhas e total. Texto à esquerda; número à direita."""
    e = _Estilos

    def escrever(linha: int, valores: list, fonte, fundo=None, borda=e.borda) -> None:  # noqa: ANN001
        for coluna, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=coluna, value=valor)
            celula.font = fonte
            celula.border = borda
            if fundo is not None:
                celula.fill = fundo
            # só a primeira coluna é texto (cliente ou item); as demais são números
            celula.alignment = e.esquerda if coluna == 1 else e.direita
            if isinstance(valor, float):
                celula.number_format = FORMATO_MOEDA

    escrever(linha_inicial, titulos, e.cabecalho, e.fundo_cabecalho)
    for deslocamento, valores in enumerate(linhas, start=1):
        escrever(linha_inicial + deslocamento, valores, e.dados)
    escrever(linha_inicial + len(linhas) + 1, total, e.totais, e.fundo_totais, e.borda_totais)

    for coluna in range(1, len(titulos) + 1):
        letra = get_column_letter(coluna)
        maior = max(
            len(_texto_da_celula(ws.cell(row=r, column=coluna).value))
            for r in range(linha_inicial, linha_inicial + len(linhas) + 2)
        )
        ws.column_dimensions[letra].width = max(maior + 4, 12)
    ws.freeze_panes = ws.cell(row=linha_inicial + 1, column=1)


def _texto_da_celula(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float):
        return f"R$ {valor:,.2f}"
    return str(valor)


def gerar_excel_geral(relatorio: RelatorioGeral) -> bytes:
    """Planilha do relatório geral: aba "Resumo" e uma aba por cliente.

    Os números são os da estrutura já calculada; nada é somado aqui além de
    transcrever. Dinheiro vai como número formatado em R$ (como no relatório por
    cliente), para o operador poder somar e conferir no próprio Excel.
    """
    e = _Estilos
    wb = Workbook()
    periodo = (
        f"Período: {relatorio.inicio.strftime('%d/%m/%Y')} a {relatorio.fim.strftime('%d/%m/%Y')}"
    )

    resumo = wb.active
    resumo.title = NOME_ABA_RESUMO
    resumo.cell(row=1, column=1, value="Todos os clientes").font = e.titulo
    resumo.cell(row=2, column=1, value=periodo).font = e.periodo
    _tabela(
        resumo,
        4,
        ["Cliente", "Peças", "Total R$"],
        [
            [secao.cliente_nome, int(secao.total_pecas), float(secao.total_valor)]
            for secao in relatorio.secoes
        ],
        ["Total geral", int(relatorio.total_pecas), float(relatorio.total_valor)],
    )

    usados = {NOME_ABA_RESUMO.casefold()}
    for secao in relatorio.secoes:
        aba = wb.create_sheet(nome_de_aba(secao.cliente_nome, usados))
        aba.cell(row=1, column=1, value=f"Cliente: {secao.cliente_nome}").font = e.titulo
        aba.cell(row=2, column=1, value=periodo).font = e.periodo
        _tabela(
            aba,
            4,
            ["Item", "Por peça", "Peças", "Subtotal"],
            [
                [
                    linha.item_nome,
                    float(linha.valor_unitario),
                    int(linha.quantidade),
                    float(linha.subtotal),
                ]
                for linha in secao.linhas
            ],
            ["Total", None, int(secao.total_pecas), float(secao.total_valor)],
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
