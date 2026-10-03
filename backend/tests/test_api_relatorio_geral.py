"""Relatório geral (todos os clientes) de ponta a ponta, contra Postgres (v1.1, Req 5).

Aqui está sob teste a consulta agrupada de verdade (``GROUP BY`` por cliente, item e
valor congelado) e a paridade entre tela e Excel. A montagem das seções sem banco
está em ``test_relatorio.py``.
"""

import io
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.core.erros import CodigoErro
from tests.conftest import criar_item_api, definir_preco

CAMINHO = "/api/relatorio/geral"
SETEMBRO = {"inicio": "2026-09-01", "fim": "2026-09-30"}


def criar_cliente(api: TestClient, nome: str) -> str:
    return api.post("/api/clientes", json={"nome": nome}).json()["id"]


def lancar(api: TestClient, cliente_id: str, data: str, linhas: dict[str, int]) -> None:
    resposta = api.post(
        "/api/lancamentos",
        json={
            "cliente_id": cliente_id,
            "data": data,
            "linhas": [{"item_id": i, "quantidade": q} for i, q in linhas.items()],
        },
    )
    assert resposta.status_code == 201, resposta.json()


@pytest.fixture
def cenario(api: TestClient) -> dict:
    """Três clientes: um com troca de preço no meio do mês, um inativo com
    pedido, um sem pedido no período."""
    aurora = criar_cliente(api, "Hotel Aurora")
    lencol = criar_item_api(api, aurora, "Lençol", "4.50", em="2026-06-01").json()["id"]
    fronha = criar_item_api(api, aurora, "Fronha", "3.50", em="2026-06-01").json()["id"]
    definir_preco(api, lencol, "4.80", em="2026-09-15")
    lancar(api, aurora, "2026-09-10", {lencol: 40, fronha: 10})
    lancar(api, aurora, "2026-09-20", {lencol: 20})
    lancar(api, aurora, "2026-10-01", {lencol: 99})  # fora do período

    pousada = criar_cliente(api, "Pousada Vista Verde")
    toalha = criar_item_api(api, pousada, "Toalha", "5.80", em="2026-06-01").json()["id"]
    lancar(api, pousada, "2026-09-05", {toalha: 12})
    api.post(f"/api/clientes/{pousada}/inativar")

    clinica = criar_cliente(api, "Clínica São Lucas")
    avental = criar_item_api(api, clinica, "Avental", "4.00", em="2026-06-01").json()["id"]
    lancar(api, clinica, "2026-08-28", {avental: 5})  # só fora do período

    return {"aurora": aurora, "pousada": pousada, "clinica": clinica}


class TestConsultaAgrupada:
    def test_secoes_linhas_e_totais(self, api: TestClient, cenario: dict) -> None:
        corpo = api.get(CAMINHO, params=SETEMBRO).json()

        assert corpo["periodo"] == {"inicio": "2026-09-01", "fim": "2026-09-30"}
        # sem pedido no período: Clínica não aparece; Pousada inativa aparece
        assert [s["cliente"]["nome"] for s in corpo["secoes"]] == [
            "Hotel Aurora",
            "Pousada Vista Verde",
        ]
        aurora = corpo["secoes"][0]
        # Lençol a dois valores (troca em 15/09) = duas linhas
        assert [
            (linha["item_nome"], linha["valor_unitario"], linha["quantidade"], linha["subtotal"])
            for linha in aurora["linhas"]
        ] == [
            ("Fronha", "3.50", 10, "35.00"),
            ("Lençol", "4.50", 40, "180.00"),
            ("Lençol", "4.80", 20, "96.00"),
        ]
        assert (aurora["total_pecas"], aurora["total_valor"]) == (70, "311.00")
        assert (corpo["total_pecas"], corpo["total_valor"]) == (82, "380.60")

    def test_total_geral_e_a_soma_dos_relatorios_por_cliente(
        self, api: TestClient, cenario: dict
    ) -> None:
        """Prova cruzada (design §5.6): duas consultas diferentes, mesmo total."""
        geral = api.get(CAMINHO, params=SETEMBRO).json()

        pecas, valor = 0, Decimal("0")
        for cliente_id in cenario.values():
            totais = api.get(
                "/api/relatorio", params={"cliente_id": cliente_id, **SETEMBRO}
            ).json()["totais"]
            pecas += totais["total_pecas"]
            valor += Decimal(totais["total_valor"])

        assert geral["total_pecas"] == pecas
        assert Decimal(geral["total_valor"]) == valor

    def test_periodo_sem_pedidos(self, api: TestClient, cenario: dict) -> None:
        corpo = api.get(CAMINHO, params={"inicio": "2025-01-01", "fim": "2025-01-31"}).json()

        assert corpo == {
            "periodo": {"inicio": "2025-01-01", "fim": "2025-01-31"},
            "secoes": [],
            "total_pecas": 0,
            "total_valor": "0.00",
        }

    def test_periodo_invertido(self, api: TestClient) -> None:
        resposta = api.get(CAMINHO, params={"inicio": "2026-09-30", "fim": "2026-09-01"})

        assert resposta.status_code == 422
        assert resposta.json()["erro"]["codigo"] == CodigoErro.PERIODO_INVALIDO.value

    def test_relatorio_por_cliente_traz_resumo_por_item(
        self, api: TestClient, cenario: dict
    ) -> None:
        corpo = api.get(
            "/api/relatorio", params={"cliente_id": cenario["aurora"], **SETEMBRO}
        ).json()

        assert [
            (linha["item_nome"], linha["valor_unitario"]) for linha in corpo["resumo_por_item"]
        ] == [
            ("Fronha", "3.50"),
            ("Lençol", "4.50"),
            ("Lençol", "4.80"),
        ]


class TestExcelGeral:
    def test_paridade_com_a_tela(self, api: TestClient, cenario: dict) -> None:
        tela = api.get(CAMINHO, params=SETEMBRO).json()

        resposta = api.get(f"{CAMINHO}/excel", params=SETEMBRO)

        assert resposta.status_code == 200
        assert (
            'filename="fechamento-todos-os-clientes-2026-09-01-a-2026-09-30.xlsx"'
            in resposta.headers["content-disposition"]
        )
        wb = openpyxl.load_workbook(io.BytesIO(resposta.content))
        assert wb.sheetnames == ["Resumo", "Hotel Aurora", "Pousada Vista Verde"]

        resumo = list(wb["Resumo"].iter_rows(min_row=5, values_only=True))
        assert resumo[-1][0] == "Total geral"
        assert resumo[-1][1] == tela["total_pecas"]
        assert Decimal(str(resumo[-1][2])) == Decimal(tela["total_valor"])
        for linha, secao in zip(resumo[:-1], tela["secoes"], strict=True):
            assert linha[0] == secao["cliente"]["nome"]
            assert linha[1] == secao["total_pecas"]
            assert Decimal(str(linha[2])) == Decimal(secao["total_valor"])

        aurora = list(wb["Hotel Aurora"].iter_rows(min_row=5, values_only=True))
        assert [(linha[0], linha[2]) for linha in aurora[:-1]] == [
            (linha["item_nome"], linha["quantidade"]) for linha in tela["secoes"][0]["linhas"]
        ]
        assert Decimal(str(aurora[-1][3])) == Decimal(tela["secoes"][0]["total_valor"])


def _moeda_br(valor: str) -> str:
    """ "380.60" da API → "R$ 380,60" como no PDF (valores do cenário < 1.000)."""
    return f"R$ {valor.replace('.', ',')}"


class TestPdfGeral:
    def test_paridade_com_a_tela(self, api: TestClient, cenario: dict) -> None:
        tela = api.get(CAMINHO, params=SETEMBRO).json()

        resposta = api.get(f"{CAMINHO}/pdf", params=SETEMBRO)

        assert resposta.status_code == 200
        assert resposta.headers["content-type"] == "application/pdf"
        assert (
            'filename="fechamento-todos-os-clientes-2026-09-01-a-2026-09-30.pdf"'
            in resposta.headers["content-disposition"]
        )
        texto = "\n".join(
            p.extract_text() or "" for p in PdfReader(io.BytesIO(resposta.content)).pages
        )
        # cada cliente da tela, na mesma ordem, com o total dele
        posicoes = [texto.index(s["cliente"]["nome"]) for s in tela["secoes"]]
        assert posicoes == sorted(posicoes)
        for secao in tela["secoes"]:
            assert _moeda_br(secao["total_valor"]) in texto
            for linha in secao["linhas"]:
                assert linha["item_nome"] in texto
                assert _moeda_br(linha["valor_unitario"]) in texto
                assert _moeda_br(linha["subtotal"]) in texto
        assert "Clínica São Lucas" not in texto  # sem pedido no período
        assert _moeda_br(tela["total_valor"]) in texto
        assert "Total geral" in texto
